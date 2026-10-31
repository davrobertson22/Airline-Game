"""Build normalized_japan_mlit.csv (monthly segment, both directions) and
normalized_japan_mlit_flow.csv (annual airport-to-airport flow, one-way) from MLIT 航空輸送統計年報 (e-Stat)."""
import sys, glob, csv, json, re, collections, warnings
warnings.filterwarnings('ignore')
sys.path.insert(0, '/home/claude/v3data/asia_uk/scripts')
from oa import *
import openpyxl
BASE = '/home/claude/v3data/asia_uk/'; RAW = BASE + 'raw/japan/'
J = {'三宅島':'MYE','三沢':'MSJ','下地島':'SHI','与論':'RNJ','与那国':'OGN','丘珠':'OKD','中標津':'SHB','中部':'NGO','久米島':'UEO','仙台':'SDJ','但馬':'TJH','佐賀':'HSG','八丈島':'HAC','出雲':'IZO','函館':'HKD','利尻':'RIS','北九州':'KKJ','北大東島':'KTD','南大東島':'MMD','南紀白浜':'SHM','名古屋':'NKM','喜界島':'KKX','壱岐':'IKI','多良間':'TRA','大分':'OIT','大島':'OIM','大阪':'ITM','大館':'ONJ','天草':'AXJ','奄美':'ASJ','奥尻':'OIR','女満別':'MMB','宮古':'MMY','宮崎':'KMI','富山':'TOY','対馬':'TSJ','屋久島':'KUM','山口宇部':'UBJ','山形':'GAJ','岡山':'OKJ','岩国':'IWK','帯広':'OBO','広島':'HIJ','庄内':'SYO','徳之島':'TKN','徳島':'TKS','成田':'NRT','新千歳':'CTS','新潟':'KIJ','旭川':'AKJ','松山':'MYJ','松本':'MMJ','沖永良部':'OKE','熊本':'KMJ','石垣':'ISG','石見':'IWJ','神戸':'UKB','福岡':'FUK','福島':'FKS','福江':'FUJ','秋田':'AXT','稚内':'WKJ','種子島':'TNE','米子':'YGJ','粟国':'AGJ','紋別':'MBE','羽田':'HND','能登':'NTQ','花巻':'HNA','茨城':'IBR','那覇':'OKA','金沢':'KMQ','釧路':'KUH','長崎':'NGS','関西':'KIX','隠岐':'OKI','青森':'AOJ','静岡':'FSZ','高松':'TAK','高知':'KCZ','鳥取':'TTJ','鹿児島':'KOJ',
     # aliases used in the 第9表 flow matrix (official airport names)
     '宮古島':'MMY','小松':'KMQ','百里':'IBR','美保':'YGJ',
     # small Izu-islands airfields / heliports without IATA codes in OurAirports
     '調布':'','新島':'','神津島':'','利島':'','御蔵島':'','青ヶ島':''}
ICAO_NOIATA = {'調布':'RJTF','新島':'RJAN','神津島':'RJAZ','利島':'(heliport)','御蔵島':'(heliport)','青ヶ島':'(heliport)'}
oa = load_oa(); IATA = by_iata(oa); G = game()
def pn(v):
    if isinstance(v, (int, float)): return int(v)
    return 0  # '-' or blank = no traffic
def clean(s): return re.sub(r'[\s　]+', '', str(s or ''))
unm = collections.Counter()
def code(nm):
    nm = clean(nm)
    if nm not in J: raise KeyError(nm)
    c = J[nm]
    return c or 'UNMAPPED:' + nm, nm
seg_rows = []; dir_tot = collections.Counter(); checks = {}
for f in sorted(glob.glob(RAW + 'jp_*.xlsx')):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    if not any('第3表' in s for s in wb.sheetnames): continue
    y = int(re.search(r'jp_(\d{4})_', f).group(1))
    for sn in wb.sheetnames:
        rows = list(wb[sn].iter_rows(values_only=True))
        state = 'await'; seg = None; months = {}
        for r in rows:
            c0 = r[0]
            if state == 'await' and c0 == 1 and isinstance(r[1], str) and '―' in r[1]:
                seg = clean(r[1]); months = {1: pn(r[5])}; state = 'cal'
            elif state == 'cal' and isinstance(c0, int) and 2 <= c0 <= 12:
                months[c0] = pn(r[5])
            elif state == 'cal' and c0 == '暦':
                state = 'caltot'; caltot = [(clean(r[1]), pn(r[5]))]
            elif state == 'caltot' and c0 == '年':
                caltot.append((clean(r[1]), pn(r[5])))
            elif state == 'caltot' and c0 == '計':
                tot = pn(r[5])
                ms = sum(v or 0 for v in months.values())
                assert len(months) == 12 and ms == tot, (f, sn, seg, ms, tot)
                assert caltot[0][0] == seg and sum(v or 0 for _, v in caltot) == tot, (seg, caltot, tot)
                a, b = [x for x in seg.split('―')]
                (oc, on), (dc, dn) = code(a), code(b)
                for mo, v in sorted(months.items()):
                    if v: seg_rows.append(['JP_MLIT_DOM_SEG', oc, dc, y, mo, int(v), 'segment', 'both'])
                    if v and (oc.startswith('UN') or dc.startswith('UN')): unm[seg] += int(v)
                for s, v in caltot:
                    x, z = s.split('―'); dir_tot[(y, code(x)[0], code(z)[0])] += int(v or 0)
                state = 'fy'
            elif state == 'fy' and c0 == '度':
                state = 'fytot'
            elif state == 'fytot' and c0 == '計':
                state = 'await'
        assert state == 'await', (f, sn, state)
    checks[f'segment CY{y}'] = sum(r[5] for r in seg_rows if r[3] == y)
with open(BASE + 'normalized_japan_mlit.csv', 'w', newline='') as fh:
    w = csv.writer(fh); w.writerow(['source','origin_iata','dest_iata','year','month','passengers','measure','direction']); w.writerows(seg_rows)
# airport-to-airport flow table (第9表, calendar year, directional)
flow = []
for f in sorted(glob.glob(RAW + 'jp_*.xlsx')):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    if not any('第9表' in s for s in wb.sheetnames): continue
    y = int(re.search(r'jp_(\d{4})_', f).group(1))
    for sn in wb.sheetnames:
        rows = list(wb[sn].iter_rows(values_only=True))
        hdr_i = next(i for i, r in enumerate(rows) if r[0] and '発空港' in str(r[0]))
        hdr = [clean(x) for x in rows[hdr_i]]
        for r in rows[hdr_i + 1:]:
            o = clean(r[0])
            if not o or o in ('合計',) or '発空港' in o: continue
            if o not in J: print('UNKNOWN ROW LABEL', y, o); continue
            for j in range(1, len(r)):
                d = hdr[j] if j < len(hdr) else ''
                if d in J and isinstance(r[j], (int, float)) and r[j] > 0:
                    flow.append(['JP_MLIT_DOM_FLOW', code(o)[0], code(d)[0], y, 0, int(r[j]), 'segment', 'one-way'])
        extra = sorted({h for h in hdr[1:] if h and h not in J})
        print('flow', y, 'non-airport header cols:', extra)
    checks[f'flow CY{y}'] = sum(r[5] for r in flow if r[3] == y)
with open(BASE + 'normalized_japan_mlit_flow.csv', 'w', newline='') as fh:
    w = csv.writer(fh); w.writerow(['source','origin_iata','dest_iata','year','month','passengers','measure','direction']); w.writerows(flow)
with open(BASE + 'mappings/japan_mlit_airport_map.csv', 'w', newline='') as fh:
    w = csv.writer(fh); w.writerow(['mlit_name','iata','oa_name','in_game','note'])
    for k, c in sorted(J.items(), key=lambda x: x[1] or 'ZZZ'):
        w.writerow([k, c, IATA.get(c, {}).get('name', '') if c else '', c in G, ('alias used in flow table' if k in ('宮古島','小松','百里','美保') else '') if c else 'no IATA code; ICAO ' + ICAO_NOIATA[k]])
json.dump(unm.most_common(), open(BASE + 'mappings/japan_unmapped.json', 'w'), ensure_ascii=False, indent=0)
print(len(seg_rows), len(flow)); print(checks); print('unmapped', unm.most_common()); print('bad', [c for c in J.values() if c and c not in IATA])
